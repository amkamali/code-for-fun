import argparse
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from pathlib import Path
import sys

# =============================================================================
# Private API Helpers
# =============================================================================
# WARNING: The following helper functions use python-docx private/internal APIs
# (_element, _p, _parent) which are not part of the official public API.
# These may break in future versions of python-docx.
# Tested with python-docx >=1.1.0,<2.0.0
# See: https://python-docx.readthedocs.io/
# =============================================================================


def _remove_paragraph(paragraph) -> None:
    """
    Remove a paragraph from a DOCX document.

    WARNING: Uses private API (_element) - may break in future python-docx versions.

    Args:
        paragraph: A python-docx Paragraph object to remove.
    """
    paragraph._element.getparent().remove(paragraph._element)


def _insert_paragraph_after(anchor, new_element) -> Paragraph:
    """
    Insert a new paragraph element after an anchor paragraph.

    WARNING: Uses private API (_p, _parent) - may break in future python-docx versions.

    Args:
        anchor: The Paragraph object after which to insert.
        new_element: The new OxmlElement to insert.

    Returns:
        A new Paragraph object wrapping the inserted element.
    """
    anchor._p.addnext(new_element)
    return Paragraph(new_element, anchor._parent)


def rewrite_profile_section(docx_path: Path, new_texts: list[str], output_path: Path):
    """
    Replaces the paragraphs in profile section of a DOCX file,
    preserving the original formatting of those paragraphs.

    Args:
        docx_path (Path): Path to the input DOCX file.
        new_texts (list[str]): A list of the new texts to replace the paragraphs
            between profile and the next heading 1.
        output_path (Path): Path to save the modified DOCX file.

    Returns:
        None
    """
    doc = Document(docx_path)

    def _copy_font(dst_font, src_font):
        if src_font is None:
            return
        try:
            dst_font.name = src_font.name
            dst_font.size = src_font.size
            dst_font.bold = src_font.bold
            dst_font.italic = src_font.italic
            dst_font.underline = src_font.underline
            if getattr(src_font, "color", None) is not None:
                dst_font.color.rgb = src_font.color.rgb
        except (AttributeError, TypeError):
            # Best-effort copy; ignore attributes not set or incompatible types
            pass

    for i, paragraph in enumerate(doc.paragraphs):
        if (
            paragraph.style.name == "Heading 1"
            and paragraph.text.lower().strip() == "profile"
        ):
            # Determine the next Heading 1 that marks the end of the Profile section
            j = None
            for idx, next_paragraph in enumerate(doc.paragraphs[i + 1 :], start=i + 1):
                if next_paragraph.style.name == "Heading 1":
                    j = idx
                    break
            if j is None:
                # No end heading found; nothing to rewrite
                break

            # Preserve style and run formatting from the first body paragraph after "Profile"
            sample_paragraph = (
                doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else None
            )
            sample_style = (
                sample_paragraph.style if sample_paragraph is not None else None
            )
            sample_run_font = (
                sample_paragraph.runs[0].font
                if (sample_paragraph is not None and sample_paragraph.runs)
                else None
            )

            # Delete paragraphs between Profile and the next Heading 1 (exclusive)
            for idx in range(j - 1, i, -1):
                p = doc.paragraphs[idx]
                _remove_paragraph(p)

            # Insert new paragraphs directly after the Profile heading
            anchor = doc.paragraphs[i]  # Profile heading paragraph
            insert_after = anchor
            for text in new_texts:
                new_p = OxmlElement("w:p")
                new_para = _insert_paragraph_after(insert_after, new_p)
                if sample_style is not None:
                    new_para.style = sample_style
                run = new_para.add_run(text)
                _copy_font(run.font, sample_run_font)
                insert_after = new_para

            doc.save(output_path)
            return

    print(
        "Profile section not found. No changes made. Output file not created.",
        file=sys.stderr,
    )
    return


def main():
    """
    Command-line interface to replace paragraphs in the Profile section of a DOCX file.
    Other sections remain unchanged, for now. Modification of other sections of the text will be added later.
    The modified document is saved to a new file.

    CLI Arguments:
        --docx_path: Path to the input DOCX file (default: ../sample_document.docx)
        --new_text_path: Path to the text file containing new text (required)
        --output_path: Path to save the modified DOCX file (default: ../new_sample_document.docx)

    Returns:
        None
    """
    parser = argparse.ArgumentParser(
        description="Replace profile section in a DOCX file."
    )
    parser.add_argument(
        "--docx_path",
        type=Path,
        default="../sample_document.docx",
        help="Path to the input DOCX file",
    )
    parser.add_argument(
        "--new_text_path",
        type=Path,
        required=True,
        help="Path to the text file containing new text for profile section",
    )
    parser.add_argument(
        "--output_path",
        type=Path,
        default="../new_sample_document.docx",
        help="Path to save the modified DOCX file",
    )
    args = parser.parse_args()

    try:
        with open(args.new_text_path, "r") as f:
            new_text = f.read().strip().splitlines()

        rewrite_profile_section(
            docx_path=args.docx_path, new_texts=new_text, output_path=args.output_path
        )
    except FileNotFoundError as e:
        print(f"Error: File not found - {e.filename}", file=sys.stderr)
        sys.exit(1)
    except PermissionError as e:
        print(f"Error: Permission denied - {e.filename}", file=sys.stderr)
        sys.exit(1)
    except OSError as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
